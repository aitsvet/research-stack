#!/usr/bin/env python3
"""PowerPoint decks (.pptx) -> paginated Markdown, one page per slide.

Same pagination as the PDF/OCR converters (`{N}` and 48 dashes before each
page, N counted from 1), so a deck reads, greps and cites like any other
corpus text. Every text frame is taken in shape order, groups are walked,
tables become pipe tables, and speaker notes follow the slide body.

Output goes next to each source as `<name>.md` unless `--out` is given, and an
existing target is left alone unless `--force` is passed.

Usage:
    pptx_md.py <file-or-dir> [...] [--out DIR] [--force] [--dry-run]
"""
import argparse
import os
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

SEP = "-" * 48


def cell(text):
    return " ".join(text.split()).replace("|", "\\|")


def blocks(shapes):
    for sh in shapes:
        try:
            kind = sh.shape_type
        except NotImplementedError:  # e.g. content parts; text frames below still apply
            kind = None
        if kind == MSO_SHAPE_TYPE.GROUP:
            yield from blocks(sh.shapes)
            continue
        if sh.has_text_frame and sh.text_frame.text.strip():
            yield sh.text_frame.text.strip().replace("\v", "\n")
        if getattr(sh, "has_table", False) and sh.has_table:
            rows = [[cell(c.text) for c in r.cells] for r in sh.table.rows]
            if rows:
                out = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * len(rows[0])]
                out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
                yield "\n".join(out)


def convert(path):
    pages = []
    for n, slide in enumerate(Presentation(path).slides, 1):
        body = "\n\n".join(blocks(slide.shapes))
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame  # None when the notes page has no body
            notes = frame.text.strip() if frame is not None else ""
            if notes:
                body += "\n\nSpeaker notes:\n\n" + notes
        pages.append(f"{{{n}}}{SEP}\n\n{body}\n")
    title = os.path.splitext(os.path.basename(path))[0]
    return f"# {title}\n\n" + "\n".join(pages)


def sources(args):
    for p in args:
        if os.path.isdir(p):
            for root, _, files in os.walk(p):
                for f in sorted(files):
                    if f.lower().endswith(".pptx") and not f.startswith("~$"):
                        yield os.path.join(root, f)
        else:
            yield p


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("paths", nargs="+")
    ap.add_argument("--out", help="output directory (default: next to the source)")
    ap.add_argument("--force", action="store_true", help="overwrite an existing .md")
    ap.add_argument("--dry-run", action="store_true", help="list what would be written")
    a = ap.parse_args()
    for src in sources(a.paths):
        dst = os.path.join(a.out or os.path.dirname(src), os.path.splitext(os.path.basename(src))[0] + ".md")
        if os.path.exists(dst) and not a.force:
            print(f"skip (exists): {dst}", file=sys.stderr)
            continue
        if a.dry_run:
            print(f"{src} -> {dst}")
            continue
        text = convert(src)
        with open(dst, "w", encoding="utf-8") as f:
            f.write(text)
        print(f"{dst}: {text.count(SEP)} slides, {len(text)} chars")


if __name__ == "__main__":
    main()
