#!/usr/bin/env python3
"""Emit a native .pptx from browser-harvested geometry (harvest.py JSON).

Positions come from a real layout engine, so they are measured, not guessed.
Text is emitted as native runs with anchor=MIDDLE — small metric differences
between Blink and PowerPoint get absorbed by the anchor instead of shifting
the block. px -> EMU is exact (9525 EMU/px @96dpi), font px -> pt is *0.75.

  emit_pptx.py geometry.json out.pptx [--asset-root DIR]
"""
import argparse
import json
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

EMU_PER_PX = 9525
ALIGN = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}


def px(v):
    return Emu(int(round(v * EMU_PER_PX)))


def rgb(css):
    m = re.match(r"rgba?\((\d+),\s*(\d+),\s*(\d+)(?:,\s*([\d.]+))?\)", css or "")
    if not m:
        return None
    if m.group(4) is not None and float(m.group(4)) == 0:
        return None  # transparent
    return RGBColor(*(int(m.group(i)) for i in (1, 2, 3)))


def add_text(slide, el, font_name):
    box = slide.shapes.add_textbox(px(el["x"]), px(el["y"]), px(el["w"]), px(el["h"]))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for m in "margin_left margin_right margin_top margin_bottom".split():
        setattr(tf, m, Emu(0))
    pt = round(el["size"] * 0.75 * 2) / 2  # clean half-point sizes
    bold = str(el["weight"]) in ("bold", "600", "700", "800", "900")
    for i, line in enumerate(el["text"].split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIGN.get(el["align"], PP_ALIGN.LEFT)
        if el.get("lineHeight"):
            p.line_spacing = round(el["lineHeight"], 2)
        r = p.add_run()
        r.text = line
        r.font.name = font_name
        r.font.size = Pt(pt)
        r.font.bold = bold
        color = rgb(el["color"])
        if color is not None:
            r.font.color.rgb = color


def add_card(slide, el):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if el["radius"] else MSO_SHAPE.RECTANGLE,
        px(el["x"]), px(el["y"]), px(el["w"]), px(el["h"]))
    if el["radius"]:
        shp.adjustments[0] = min(el["radius"] / min(el["w"], el["h"]), 0.5)
    fill = rgb(el["bgColor"])
    if fill is not None:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    else:
        shp.fill.background()
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("geometry")
    ap.add_argument("out")
    ap.add_argument("--asset-root", default=".", help="base dir for data-src image paths")
    ap.add_argument("--font", default="Arial", help="font name written into the pptx")
    a = ap.parse_args()

    slides = json.load(open(a.geometry, encoding="utf-8"))
    prs = Presentation()
    prs.slide_width = px(slides[0]["w"])
    prs.slide_height = px(slides[0]["h"])
    for s in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        bg = rgb(s["bg"])
        if bg is not None:
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = bg
        for el in s["els"]:
            if el["kind"] == "card":
                add_card(slide, el)
            elif el["kind"] == "image":
                path = Path(a.asset_root) / el["src"]
                slide.shapes.add_picture(str(path), px(el["x"]), px(el["y"]),
                                         px(el["w"]), px(el["h"]))
            elif el["kind"] == "text":
                add_text(slide, el, a.font)
    prs.save(a.out)
    print(f"wrote {a.out}: {len(slides)} slide(s)")


if __name__ == "__main__":
    main()
